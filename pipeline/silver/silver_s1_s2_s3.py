import os
import sys
import json
import glob
import shutil
import hashlib
from datetime import datetime, timezone

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from shapely.geometry import Polygon, MultiPolygon
from shapely.ops import unary_union


# ============================================================
# EpiScope - S1/S2/S3 Silver Transformation
# ============================================================

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))

RAW_DIR = os.path.join(BASE_DIR, "data", "raw")
SILVER_DIR = os.path.join(BASE_DIR, "data", "silver")
QUARANTINE_DIR = os.path.join(BASE_DIR, "data", "quarantine")
QUALITY_DIR = os.path.join(BASE_DIR, "data", "gold", "quality")

S1_RAW = os.path.join(
    RAW_DIR,
    "S1_boundaries",
    "bangladesh_districts.geojson",
)

S2_RAW = os.path.join(
    RAW_DIR,
    "S2_population",
    "bangladesh_bbs_population-and-housing-census-dataset_2022_admin-02.xlsx",
)

S3_RAW_DIR = os.path.join(RAW_DIR, "S3_weather")

S1_SILVER = os.path.join(SILVER_DIR, "S1_boundaries")
S2_SILVER = os.path.join(SILVER_DIR, "S2_population")
S3_SILVER = os.path.join(SILVER_DIR, "S3_weather")

S1_QUARANTINE = os.path.join(QUARANTINE_DIR, "S1_boundaries")
S2_QUARANTINE = os.path.join(QUARANTINE_DIR, "S2_population")
S3_QUARANTINE = os.path.join(QUARANTINE_DIR, "S3_weather")

S1_QUALITY = os.path.join(QUALITY_DIR, "s1_boundaries")
S2_QUALITY = os.path.join(QUALITY_DIR, "s2_population")
S3_QUALITY = os.path.join(QUALITY_DIR, "s3_weather")

GENERATED_AT = datetime.now(timezone.utc).isoformat()


# ============================================================
# Helpers
# ============================================================

def remove_output(path):
    if os.path.exists(path):
        print(f"Removing old output: {path}")
        shutil.rmtree(path)


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)


def write_parquet(df, path):
    ensure_dir(os.path.dirname(path))
    table = pa.Table.from_pandas(df, preserve_index=False)
    pq.write_table(table, path, compression="snappy")


def write_quality_report(dataset_name, quality_dir, summary):
    remove_output(quality_dir)
    ensure_dir(quality_dir)

    quality_df = pd.DataFrame(
        [
            {
                "dataset": dataset_name,
                "total_rows": summary["total_rows"],
                "clean_rows": summary["clean_rows"],
                "quarantine_rows": summary["quarantine_rows"],
                "generated_at": summary["generated_at"],
            }
        ]
    )

    for rule, count in summary["quality_rules"].items():
        quality_df[rule] = int(count)

    write_parquet(
        quality_df,
        os.path.join(quality_dir, "part-00000.parquet"),
    )

    with open(
        os.path.join(
            os.path.dirname(quality_dir),
            f"{dataset_name}_summary.json",
        ),
        "w",
        encoding="utf-8",
    ) as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)

    with open(
        os.path.join(quality_dir, "_SUCCESS"),
        "w",
        encoding="utf-8",
    ) as f:
        f.write("")


def dataframe_checksum(df):
    """
    Deterministic checksum for the Silver transformation output.
    """
    if df.empty:
        return hashlib.sha256(b"EMPTY").hexdigest()

    normalized = df.fillna("").astype(str)

    csv_bytes = normalized.to_csv(
        index=False,
        lineterminator="\n",
    ).encode("utf-8")

    return hashlib.sha256(csv_bytes).hexdigest()


# ============================================================
# S1 - Administrative Boundaries
# ============================================================

def process_s1():
    print()
    print("=" * 70)
    print("S1 - ADMINISTRATIVE BOUNDARIES")
    print("=" * 70)

    if not os.path.exists(S1_RAW):
        raise FileNotFoundError(f"S1 raw file not found: {S1_RAW}")

    with open(S1_RAW, "r", encoding="utf-8") as f:
        data = json.load(f)

    features = data.get("features", [])

    print(f"Raw features: {len(features)}")

    rows = []
    quarantine_rows = []

    quality_rules = {
        "missing_district_name": 0,
        "missing_district_pcode": 0,
        "missing_division_name": 0,
        "missing_geometry": 0,
        "duplicate_district_pcode": 0,
        "duplicate_district_name": 0,
    }

    for idx, feature in enumerate(features):
        attrs = feature.get("attributes") or {}
        geometry = feature.get("geometry") or {}

        district_name = attrs.get("adm2_en")
        district_pcode = attrs.get("adm2_pcode")
        division_name = attrs.get("adm1_en")
        division_pcode = attrs.get("adm1_pcode")

        rings = geometry.get("rings")

        if not district_name:
            quality_rules["missing_district_name"] += 1

        if not district_pcode:
            quality_rules["missing_district_pcode"] += 1

        if not division_name:
            quality_rules["missing_division_name"] += 1

        if not rings:
            quality_rules["missing_geometry"] += 1

        row = {
            "district_name": district_name,
            "district_pcode": district_pcode,
            "division_name": division_name,
            "division_pcode": division_pcode,
            "adm0_pcode": attrs.get("adm0_pcode"),
            "adm0_name": attrs.get("adm0_en"),
            "fid": attrs.get("fid"),
            "shape_length": attrs.get("shape_leng"),
            "shape_area": attrs.get("SHAPE__Area"),
            "geometry_rings_json": json.dumps(
                rings,
                ensure_ascii=False,
                separators=(",", ":"),
            ) if rings else None,
            "source_feature_index": idx,
        }

        # Representative point from the ArcGIS rings where possible.
        representative_lon = None
        representative_lat = None

        try:
            polygons = []

            for ring in rings or []:
                if len(ring) >= 4:
                    polygon = Polygon(ring)
                    if not polygon.is_empty and polygon.is_valid:
                        polygons.append(polygon)

            if polygons:
                geom = unary_union(polygons)

                if not geom.is_empty:
                    point = geom.representative_point()
                    representative_lon = float(point.x)
                    representative_lat = float(point.y)

        except Exception:
            representative_lon = None
            representative_lat = None

        row["representative_longitude"] = representative_lon
        row["representative_latitude"] = representative_lat

        reasons = []

        if not district_name:
            reasons.append("missing_district_name")

        if not district_pcode:
            reasons.append("missing_district_pcode")

        if not division_name:
            reasons.append("missing_division_name")

        if not rings:
            reasons.append("missing_geometry")

        row["dq_reason"] = "|".join(reasons) if reasons else None

        if reasons:
            quarantine_rows.append(row)
        else:
            rows.append(row)

    clean_df = pd.DataFrame(rows)
    quarantine_df = pd.DataFrame(quarantine_rows)

    if not clean_df.empty:
        quality_rules["duplicate_district_pcode"] = int(
            clean_df["district_pcode"].duplicated(keep=False).sum()
        )
        quality_rules["duplicate_district_name"] = int(
            clean_df["district_name"].duplicated(keep=False).sum()
        )

        duplicate_mask = (
            clean_df["district_pcode"].duplicated(keep=False)
            | clean_df["district_name"].duplicated(keep=False)
        )

        if duplicate_mask.any():
            duplicate_rows = clean_df.loc[duplicate_mask].copy()

            duplicate_rows["dq_reason"] = duplicate_rows.apply(
                lambda r: (
                    "duplicate_district_pcode"
                    if clean_df["district_pcode"]
                    .duplicated(keep=False)
                    .loc[r.name]
                    else ""
                )
                + (
                    "|duplicate_district_name"
                    if clean_df["district_name"]
                    .duplicated(keep=False)
                    .loc[r.name]
                    else ""
                ),
                axis=1,
            )

            clean_df = clean_df.loc[~duplicate_mask].copy()
            quarantine_df = pd.concat(
                [quarantine_df, duplicate_rows],
                ignore_index=True,
            )

    total_rows = len(features)
    clean_rows = len(clean_df)
    quarantine_count = len(quarantine_df)

    remove_output(S1_SILVER)
    remove_output(S1_QUARANTINE)

    ensure_dir(S1_SILVER)
    ensure_dir(S1_QUARANTINE)

    print(f"Total rows      : {total_rows}")
    print(f"Clean rows      : {clean_rows}")
    print(f"Quarantine rows : {quarantine_count}")

    if not clean_df.empty:
        write_parquet(
            clean_df,
            os.path.join(S1_SILVER, "part-00000.parquet"),
        )

    if not quarantine_df.empty:
        write_parquet(
            quarantine_df,
            os.path.join(S1_QUARANTINE, "part-00000.parquet"),
        )

    with open(
        os.path.join(S1_SILVER, "_SUCCESS"),
        "w",
        encoding="utf-8",
    ) as f:
        f.write("")

    summary = {
        "dataset": "s1_boundaries",
        "total_rows": total_rows,
        "clean_rows": clean_rows,
        "quarantine_rows": quarantine_count,
        "quality_rules": quality_rules,
        "generated_at": GENERATED_AT,
        "checksum": dataframe_checksum(clean_df),
        "notes": [
            "Original ArcGIS rings are preserved as geometry_rings_json.",
            "Representative points are derived only where polygon construction succeeds.",
            "No source geometry is modified in Bronze.",
        ],
    }

    write_quality_report(
        "s1_boundaries",
        S1_QUALITY,
        summary,
    )

    print()
    print("S1 COMPLETED")
    print(f"Silver     : {S1_SILVER}")
    print(f"Quarantine : {S1_QUARANTINE}")
    print(f"Quality    : {S1_QUALITY}")


# ============================================================
# S2 - Population
# ============================================================

def process_s2():
    print()
    print("=" * 70)
    print("S2 - DISTRICT POPULATION")
    print("=" * 70)

    if not os.path.exists(S2_RAW):
        raise FileNotFoundError(f"S2 raw file not found: {S2_RAW}")

    wb = load_workbook(
        S2_RAW,
        read_only=True,
        data_only=True,
    )

    if "Merged_All_Table" not in wb.sheetnames:
        raise ValueError(
            "Merged_All_Table sheet not found in S2 workbook."
        )

    ws = wb["Merged_All_Table"]

    rows = ws.iter_rows(values_only=True)

    headers = next(rows)

    headers = [
        str(h).strip() if h is not None else ""
        for h in headers
    ]

    records = []

    for row in rows:
        record = {}

        for i, header in enumerate(headers):
            value = row[i] if i < len(row) else None
            record[header] = value

        records.append(record)

    wb.close()

    df = pd.DataFrame(records)

    print(f"Raw rows: {len(df)}")
    print(f"Raw columns: {len(df.columns)}")

    required_columns = [
        "Division",
        "District",
        "Division_Geocode",
        "District_Geocode",
        "Household_Total",
        "Population_Total",
        "Population Density",
    ]

    missing_columns = [
        col for col in required_columns
        if col not in df.columns
    ]

    if missing_columns:
        raise ValueError(
            "Missing S2 columns: "
            + ", ".join(missing_columns)
        )

    # Keep the core population fields needed by EpiScope.
    silver_df = df[
        [
            "Division",
            "District",
            "Division_Geocode",
            "District_Geocode",
            "Household_Total",
            "Population_Total",
            "Population Density",
        ]
    ].copy()

    silver_df = silver_df.rename(
        columns={
            "Division": "division_name",
            "District": "district_name_source",
            "Division_Geocode": "division_geocode",
            "District_Geocode": "district_geocode",
            "Household_Total": "household_total",
            "Population_Total": "population_total",
            "Population Density": "population_density",
        }
    )

    # Preserve the original BBS spelling.
    # Canonical names are only for joining with S1/S4/S5.
    canonical_mapping = {
   	 "Barishal": "Barisal",
    	 "Chattogram": "Chittagong",
    	 "Brahmanbaria": "Brahamanbaria",
    	 "Bogura": "Bogra",
    	 "Jashore": "Jessore",
   	 "Cumilla": "Comilla",
         "Cox's Bazar": "Cox's Bazar",
    	 "Nawabganj": "Chapainababganj",
    	 "Maulvibazar": "Moulvibazar",
     }

    silver_df["district_name"] = (
        silver_df["district_name_source"]
        .astype(str)
        .str.strip()
    )

    silver_df["district_name_canonical"] = (
        silver_df["district_name"]
        .map(canonical_mapping)
        .fillna(silver_df["district_name"])
    )

    # Numeric conversion.
    silver_df["division_geocode"] = pd.to_numeric(
        silver_df["division_geocode"],
        errors="coerce",
    )

    silver_df["district_geocode"] = pd.to_numeric(
        silver_df["district_geocode"],
        errors="coerce",
    )

    silver_df["household_total"] = pd.to_numeric(
        silver_df["household_total"],
        errors="coerce",
    )

    silver_df["population_total"] = pd.to_numeric(
        silver_df["population_total"],
        errors="coerce",
    )

    silver_df["population_density"] = pd.to_numeric(
        silver_df["population_density"],
        errors="coerce",
    )

    quality_rules = {
        "missing_division": 0,
        "missing_district": 0,
        "missing_district_geocode": 0,
        "invalid_household_total": 0,
        "invalid_population_total": 0,
        "invalid_population_density": 0,
        "duplicate_district_geocode": 0,
        "duplicate_district_name": 0,
    }

    silver_df["dq_reason"] = ""

    def add_reason(mask, reason):
        silver_df.loc[mask, "dq_reason"] = silver_df.loc[
            mask, "dq_reason"
        ].apply(
            lambda x: reason if not x else f"{x}|{reason}"
        )

    mask = (
        silver_df["division_name"].isna()
        | (silver_df["division_name"].astype(str).str.strip() == "")
    )
    quality_rules["missing_division"] = int(mask.sum())
    add_reason(mask, "missing_division")

    mask = (
        silver_df["district_name_source"].isna()
        | (silver_df["district_name_source"].astype(str).str.strip() == "")
    )
    quality_rules["missing_district"] = int(mask.sum())
    add_reason(mask, "missing_district")

    mask = silver_df["district_geocode"].isna()
    quality_rules["missing_district_geocode"] = int(mask.sum())
    add_reason(mask, "missing_district_geocode")

    mask = (
        silver_df["household_total"].isna()
        | (silver_df["household_total"] < 0)
    )
    quality_rules["invalid_household_total"] = int(mask.sum())
    add_reason(mask, "invalid_household_total")

    mask = (
        silver_df["population_total"].isna()
        | (silver_df["population_total"] < 0)
    )
    quality_rules["invalid_population_total"] = int(mask.sum())
    add_reason(mask, "invalid_population_total")

    mask = (
        silver_df["population_density"].isna()
        | (silver_df["population_density"] < 0)
    )
    quality_rules["invalid_population_density"] = int(mask.sum())
    add_reason(mask, "invalid_population_density")

    duplicate_code = silver_df["district_geocode"].duplicated(
        keep=False
    )
    quality_rules["duplicate_district_geocode"] = int(
        duplicate_code.sum()
    )
    add_reason(
        duplicate_code,
        "duplicate_district_geocode",
    )

    duplicate_name = silver_df["district_name"].duplicated(
        keep=False
    )
    quality_rules["duplicate_district_name"] = int(
        duplicate_name.sum()
    )
    add_reason(
        duplicate_name,
        "duplicate_district_name",
    )

    silver_df["dq_reason"] = silver_df["dq_reason"].replace(
        "",
        None,
    )

    quarantine_mask = silver_df["dq_reason"].notna()

    clean_df = silver_df.loc[
        ~quarantine_mask
    ].copy()

    quarantine_df = silver_df.loc[
        quarantine_mask
    ].copy()

    remove_output(S2_SILVER)
    remove_output(S2_QUARANTINE)

    ensure_dir(S2_SILVER)
    ensure_dir(S2_QUARANTINE)

    print(f"Clean rows      : {len(clean_df)}")
    print(f"Quarantine rows : {len(quarantine_df)}")

    if not clean_df.empty:
        write_parquet(
            clean_df,
            os.path.join(S2_SILVER, "part-00000.parquet"),
        )

    if not quarantine_df.empty:
        write_parquet(
            quarantine_df,
            os.path.join(S2_QUARANTINE, "part-00000.parquet"),
        )

    with open(
        os.path.join(S2_SILVER, "_SUCCESS"),
        "w",
        encoding="utf-8",
    ) as f:
        f.write("")

    summary = {
        "dataset": "s2_population",
        "total_rows": len(silver_df),
        "clean_rows": len(clean_df),
        "quarantine_rows": len(quarantine_df),
        "quality_rules": quality_rules,
        "generated_at": GENERATED_AT,
        "checksum": dataframe_checksum(clean_df),
        "source_sheet": "Merged_All_Table",
        "notes": [
            "Population_Total is retained as the district population measure.",
            "Original BBS district spelling is retained in district_name_source.",
            "district_name_canonical is provided for cross-source joins.",
        ],
    }

    write_quality_report(
        "s2_population",
        S2_QUALITY,
        summary,
    )

    print()
    print("S2 COMPLETED")
    print(f"Silver     : {S2_SILVER}")
    print(f"Quarantine : {S2_QUARANTINE}")
    print(f"Quality    : {S2_QUALITY}")


# ============================================================
# S3 - Daily Weather
# ============================================================

def process_s3():
    print()
    print("=" * 70)
    print("S3 - DAILY WEATHER")
    print("=" * 70)

    files = sorted(
        glob.glob(
            os.path.join(
                S3_RAW_DIR,
                "*.json",
            )
        )
    )

    if not files:
        raise FileNotFoundError(
            f"No S3 weather JSON files found in {S3_RAW_DIR}"
        )

    rows = []
    quarantine_rows = []

    quality_rules = {
        "missing_district_name": 0,
        "missing_date": 0,
        "missing_t2m": 0,
        "missing_t2m_max": 0,
        "missing_t2m_min": 0,
        "missing_precipitation": 0,
        "missing_rh2m": 0,
        "invalid_date": 0,
        "duplicate_district_date": 0,
    }

    expected_parameters = [
        "T2M",
        "T2M_MAX",
        "T2M_MIN",
        "PRECTOTCORR",
        "RH2M",
    ]

    for file_path in files:
        with open(
            file_path,
            "r",
            encoding="utf-8",
        ) as f:
            data = json.load(f)

        file_name = os.path.basename(file_path)
        district_name = os.path.splitext(file_name)[0]

        parameters = (
            data.get("properties", {})
            .get("parameter", {})
        )

        all_dates = set()

        for parameter_name in expected_parameters:
            all_dates.update(
                parameters.get(
                    parameter_name,
                    {},
                ).keys()
            )

        for date_str in sorted(all_dates):
            row = {
                "district_name": district_name,
                "date": date_str,
                "t2m": parameters.get(
                    "T2M",
                    {},
                ).get(date_str),
                "t2m_max": parameters.get(
                    "T2M_MAX",
                    {},
                ).get(date_str),
                "t2m_min": parameters.get(
                    "T2M_MIN",
                    {},
                ).get(date_str),
                "precipitation_mm": parameters.get(
                    "PRECTOTCORR",
                    {},
                ).get(date_str),
                "rh2m": parameters.get(
                    "RH2M",
                    {},
                ).get(date_str),
                "source_file": file_name,
            }

            reasons = []

            if not district_name:
                reasons.append("missing_district_name")

            parsed_date = pd.to_datetime(
                date_str,
                format="%Y%m%d",
                errors="coerce",
            )

            if pd.isna(parsed_date):
                reasons.append("invalid_date")
                quality_rules["invalid_date"] += 1
            else:
                row["date"] = parsed_date.date().isoformat()

            if row["t2m"] is None:
                reasons.append("missing_t2m")

            if row["t2m_max"] is None:
                reasons.append("missing_t2m_max")

            if row["t2m_min"] is None:
                reasons.append("missing_t2m_min")

            if row["precipitation_mm"] is None:
                reasons.append("missing_precipitation")

            if row["rh2m"] is None:
                reasons.append("missing_rh2m")

            if "missing_district_name" in reasons:
                quality_rules["missing_district_name"] += 1

            if row["t2m"] is None:
                quality_rules["missing_t2m"] += 1

            if row["t2m_max"] is None:
                quality_rules["missing_t2m_max"] += 1

            if row["t2m_min"] is None:
                quality_rules["missing_t2m_min"] += 1

            if row["precipitation_mm"] is None:
                quality_rules["missing_precipitation"] += 1

            if row["rh2m"] is None:
                quality_rules["missing_rh2m"] += 1

            row["dq_reason"] = (
                "|".join(reasons)
                if reasons
                else None
            )

            if reasons:
                quarantine_rows.append(row)
            else:
                rows.append(row)

    clean_df = pd.DataFrame(rows)
    quarantine_df = pd.DataFrame(quarantine_rows)

    if not clean_df.empty:
        duplicate_mask = clean_df.duplicated(
            subset=[
                "district_name",
                "date",
            ],
            keep=False,
        )

        quality_rules["duplicate_district_date"] = int(
            duplicate_mask.sum()
        )

        if duplicate_mask.any():
            duplicate_rows = clean_df.loc[
                duplicate_mask
            ].copy()

            duplicate_rows["dq_reason"] = (
                "duplicate_district_date"
            )

            clean_df = clean_df.loc[
                ~duplicate_mask
            ].copy()

            quarantine_df = pd.concat(
                [
                    quarantine_df,
                    duplicate_rows,
                ],
                ignore_index=True,
            )

    if not clean_df.empty:
        numeric_columns = [
            "t2m",
            "t2m_max",
            "t2m_min",
            "precipitation_mm",
            "rh2m",
        ]

        for column in numeric_columns:
            clean_df[column] = pd.to_numeric(
                clean_df[column],
                errors="coerce",
            )

        clean_df["date"] = pd.to_datetime(
            clean_df["date"],
            errors="coerce",
        ).dt.date.astype(str)

    remove_output(S3_SILVER)
    remove_output(S3_QUARANTINE)

    ensure_dir(S3_SILVER)
    ensure_dir(S3_QUARANTINE)

    print(f"Weather files  : {len(files)}")
    print(f"Total clean    : {len(clean_df)}")
    print(f"Quarantine     : {len(quarantine_df)}")

    if not clean_df.empty:
        # Partition by district so later district-level analytics
        # can prune unrelated weather data.
        for district_name, district_df in clean_df.groupby(
            "district_name"
        ):
            safe_name = (
                str(district_name)
                .replace("/", "_")
                .replace("\\", "_")
                .replace(" ", "_")
            )

            district_dir = os.path.join(
                S3_SILVER,
                f"district_name={safe_name}",
            )

            write_parquet(
                district_df,
                os.path.join(
                    district_dir,
                    "part-00000.parquet",
                ),
            )

    if not quarantine_df.empty:
        write_parquet(
            quarantine_df,
            os.path.join(
                S3_QUARANTINE,
                "part-00000.parquet",
            ),
        )

    with open(
        os.path.join(S3_SILVER, "_SUCCESS"),
        "w",
        encoding="utf-8",
    ) as f:
        f.write("")

    summary = {
        "dataset": "s3_weather",
        "total_rows": len(clean_df) + len(quarantine_df),
        "clean_rows": len(clean_df),
        "quarantine_rows": len(quarantine_df),
        "quality_rules": quality_rules,
        "generated_at": GENERATED_AT,
        "checksum": dataframe_checksum(clean_df),
        "parameters": expected_parameters,
        "notes": [
            "Daily NASA POWER observations are retained.",
            "Weather Silver is partitioned by district_name.",
            "No weather values are imputed during Silver transformation.",
        ],
    }

    write_quality_report(
        "s3_weather",
        S3_QUALITY,
        summary,
    )

    print()
    print("S3 COMPLETED")
    print(f"Silver     : {S3_SILVER}")
    print(f"Quarantine : {S3_QUARANTINE}")
    print(f"Quality    : {S3_QUALITY}")


# ============================================================
# Main
# ============================================================

def main():
    print()
    print("=" * 70)
    print("EpiScope S1/S2/S3 SILVER TRANSFORMATION")
    print("=" * 70)
    print(f"Project directory: {BASE_DIR}")
    print(f"Generated at     : {GENERATED_AT}")

    process_s1()
    process_s2()
    process_s3()

    print()
    print("=" * 70)
    print("ALL S1/S2/S3 SILVER TRANSFORMATIONS COMPLETED")
    print("=" * 70)


if __name__ == "__main__":
    main()